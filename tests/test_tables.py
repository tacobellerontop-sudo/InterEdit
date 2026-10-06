import os
_tmp_config = __import__('tempfile').mkdtemp(prefix='smoothcursor-test-')
os.environ['XDG_CONFIG_HOME'] = os.path.join(_tmp_config, 'config')
os.environ['XDG_DATA_HOME'] = os.path.join(_tmp_config, 'data')
import re, sys, tempfile
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from PySide6.QtWidgets import QApplication
from PySide6.QtGui import QTextCursor, QTextDocument
app = QApplication([])
from smoothcursor.views.window import (build_markdown_table, image_markdown, text_to_docx,
                    docx_to_text, docx_to_document, MainWindow)
from smoothcursor.editor.editor import CodeEditor
from smoothcursor.core.settings import Settings, DEFAULTS
from smoothcursor.core.theme import THEMES
from PIL import Image
import docx as pydocx

fails = []
def check(name, got, exp):
    ok = got == exp
    print(('PASS ' if ok else 'FAIL ') + name +
          ('' if ok else ' got=%r exp=%r' % (got, exp)))
    if not ok:
        fails.append(name)

# builders
t, off = build_markdown_table(2, 3, True)
check('table builder', t, '| Header 1 | Header 2 | Header 3 |\n'
      '|----------|----------|----------|\n|  |  |  |\n|  |  |  |')
check('table offset', (t[:off].endswith('| '), t[off]), (True, ' '))
check('no-header', build_markdown_table(1, 2, False), ('|  |  |', 2))
check('img relative', image_markdown('/a/b/pic.png', '/a/b/doc.docx'),
      '![pic](pic.png)')
check('img absolute', image_markdown('/x/pic.png', '/a/b/doc.docx'),
      '![pic](/x/pic.png)')

# editor helpers
def make(text):
    s = Settings()
    for k, v in DEFAULTS.items():
        s.set(k, v)
    ed = CodeEditor(s, THEMES['dark'])
    ed.setPlainText(text)
    return ed
ed = make('| a | b |\n|---|---|\n| 1 | 22 |')
ed.textCursor().setPosition(0)
c = ed.textCursor(); c.setPosition(0); ed.setTextCursor(c)
check('format table', ed.format_markdown_table(), True)
check('formatted', ed.toPlainText(), '| a | b  |\n|---|----|\n| 1 | 22 |')
check('format noop', make('plain').format_markdown_table(), False)
ed = make('ab')
ed.insert_snippet('XY', 1)
check('snippet', (ed.toPlainText(), ed.textCursor().position()),
      ('XYab', 1))

# export (marker path)
tmp = tempfile.mkdtemp()
img = os.path.join(tmp, 'pic.png')
Image.new('RGB', (800, 600), 'red').save(img)
out = os.path.join(tmp, 'out.docx')
text_to_docx(out, '# T\n\n| Name | Age |\n|------|-----|\n| Ann | 30 |\n\n'
                  'See ![pic](pic.png).\n\n![pic](pic.png)\n')
d = pydocx.Document(out)
check('export tables', len(d.tables), 1)
row0 = [c.text for c in d.tables[0].rows[0].cells]
check('header cells', row0, ['Name', 'Age'])
check('header bold',
      d.tables[0].rows[0].cells[0].paragraphs[0].runs[0].bold, True)
check('body', [[c.text for c in r.cells] for r in d.tables[0].rows[1:]],
      [['Ann', '30']])
check('pictures', len(d.inline_shapes), 2)
out2 = os.path.join(tmp, 'out2.docx')
text_to_docx(out2, 'See ![ghost](nope.png).')
d2 = pydocx.Document(out2)
check('missing img literal', d2.paragraphs[0].text,
      'See ![ghost](nope.png).')

# import
src = os.path.join(tmp, 'src.docx')
d0 = pydocx.Document()
d0.add_heading('Report', 1)
r = d0.add_paragraph().add_run('hello'); r.bold = True
tb = d0.add_table(rows=2, cols=2)
tb.rows[0].cells[0].text = 'A'; tb.rows[0].cells[1].text = 'B'
tb.rows[1].cells[0].text = '1'; tb.rows[1].cells[1].text = '2'
d0.add_paragraph().add_run().add_picture(img)
d0.add_paragraph('tail')
d0.save(src)
md = docx_to_text(src)
check('import heading/bold', ('# Report' in md, '**hello**' in md),
      (True, True))
check('import table', all(s in md for s in
      ('| A | B |', '|---|---|', '| 1 | 2 |')), True)
m = re.search(r'!\[image\d+\]\(([^)]+)\)', md)
check('import image ref', bool(m), True)
check('media saved', os.path.exists(m.group(1)) if m else False, True)
qdoc = QTextDocument()
docx_to_document(src, qdoc)
plain = qdoc.toPlainText()
# .docx tabs now use real QTextTables, not `|…|` markdown lines.
found = []
blk = qdoc.firstBlock()
while blk.isValid():
    probe = QTextCursor(qdoc)
    probe.setPosition(blk.position())
    tbl = probe.currentTable()
    if tbl is not None:
        try:
            anchor = tbl.cellAt(0, 0).firstCursorPosition().position()
        except Exception:
            anchor = None
        if not found or found[-1][0] != anchor:
            rows = tbl.rows()
            cols = tbl.columns()
            cells = []
            for _r in range(rows):
                row = []
                for _c in range(cols):
                    try:
                        _cell = tbl.cellAt(_r, _c)
                        _b = _cell.firstCursorPosition().block().text()
                    except Exception:
                        _b = ""
                    row.append(_b)
                cells.append(row)
            hdr_bold = False
            try:
                _f = tbl.cellAt(0, 0).firstCursorPosition().block()
                _it = _f.begin()
                if not _it.atEnd():
                    _frag = _it.fragment()
                    from PySide6.QtGui import QFont as _QF
                    hdr_bold = _frag.charFormat().fontWeight() > _QF.Normal
            except Exception:
                hdr_bold = False
            found.append((anchor, cells, hdr_bold))
    blk = blk.next()
check('qdoc real table cells',
      found[0][1] if found else None, [['A', 'B'], ['1', '2']])
check('qdoc header bold', found[0][2] if found else None, True)
check('qdoc image ref', bool(re.search(r'!\[image\d+\]', plain)), True)
check('qdoc no pipes', ('| A | B |' in plain), False)

# roundtrip
rt = os.path.join(tmp, 'rt.docx')
text_to_docx(rt, '| X | Y |\n|---|---|\n| 1 | 2 |')
md2 = docx_to_text(rt)
check('roundtrip', md2, '| X | Y |\n|---|---|\n| 1 | 2 |')

# rich editor: Insert > Table creates a real QTextTable, Tab moves cells,
# Format Table styles it, and save round-trips through real Word tables.
from smoothcursor.editor.rich import RichEditor
s = Settings()
for k, v in DEFAULTS.items():
    s.set(k, v)
red = RichEditor(s, THEMES['dark'])
red.insert_real_table(2, 3, True)
tbl = red.textCursor().currentTable()
check('rich insert table', (tbl.rows(), tbl.columns()), (3, 3))
check('rich header text',
      [tbl.cellAt(0, c).firstCursorPosition().block().text()
       for c in range(3)],
      ['Header 1', 'Header 2', 'Header 3'])
# type into first body cell, then Tab to the next cell
red.textCursor().insertText('cell!')
from PySide6.QtGui import QKeyEvent
from PySide6.QtCore import Qt as _Qt, QEvent as _QEvent
pos_before = red.textCursor().position()
red._move_to_cell(1)
check('rich tab next cell', red.textCursor().position() != pos_before, True)
check('rich format table', red.format_markdown_table(), True)
# legacy pipes convert to a real table
red2 = RichEditor(s, THEMES['dark'])
red2.setPlainText('| a | b |\n|---|---|\n| 1 | 2 |')
_cur = red2.textCursor()
_cur.setPosition(0)
red2.setTextCursor(_cur)
check('rich convert pipes', red2.format_markdown_table(), True)
check('rich converted is table',
      red2.textCursor().currentTable() is not None, True)
# export a rich document with a real table -> real Word table
rt2 = os.path.join(tmp, 'rt2.docx')
text_to_docx(rt2, '', document=red.document())
d3 = pydocx.Document(rt2)
check('rich export tables', len(d3.tables), 1)
check('rich export header',
      [c.text for c in d3.tables[0].rows[0].cells],
      ['Header 1', 'Header 2', 'Header 3'])
# re-import keeps it a real table
qdoc2 = QTextDocument()
docx_to_document(rt2, qdoc2)
blk2 = qdoc2.firstBlock()
has_table = False
while blk2.isValid():
    _p = QTextCursor(qdoc2)
    _p.setPosition(blk2.position())
    if _p.currentTable() is not None:
        has_table = True
        break
    blk2 = blk2.next()
check('rich reimport table', has_table, True)

# --- save-name: no doubled ".docx" ---
from smoothcursor.views.window import MainWindow as _MW
class _FakeEd:
    file_path = None
    untitled_name = "Untitled 1.docx"
check('save suggest single docx',
      os.path.basename(_MW.suggested_save_path(_FakeEd())),
      "Untitled 1.docx")
_FakeEd.untitled_name = "Untitled 2"
check('save suggest adds docx',
      os.path.basename(_MW.suggested_save_path(_FakeEd())),
      "Untitled 2.docx")
_FakeEd.file_path = "/tmp/keep.docx"
_FakeEd.untitled_name = ""
check('save suggest keeps path', _MW.suggested_save_path(_FakeEd()),
      "/tmp/keep.docx")
check('normalize double docx',
      _MW.normalize_save_path("/tmp/Untitled 1.docx.docx"),
      "/tmp/Untitled 1.docx")

# --- table customisation: style, rows/cols, fill round-trip ---
red3 = RichEditor(s, THEMES['dark'])
red3.insert_real_table(2, 2, True, border=0, padding=8,
                       header_fill="#c6efce", header_bold=True)
tbl3 = red3.textCursor().currentTable()
check('custom style', red3.get_table_style(),
      {'border': 0, 'padding': 8.0, 'spacing': 0.0,
       'header_fill': '#c6efce', 'header_bold': True})
red3.insert_table_row(below=True)
check('insert row', tbl3.rows(), 4)
red3.insert_table_column(right=True)
check('insert col', tbl3.columns(), 3)
red3.set_table_style(border=2, padding=6, header_fill="#d6e4f0",
                     header_bold=False)
st = red3.get_table_style()
check('restyle border/pad', (st["border"], st["padding"],
                             st["header_fill"], st["header_bold"]),
      (2, 6.0, "#d6e4f0", False))
check('toggle header', red3.toggle_header_row(), True)
check('header toggled on', red3.get_table_style()["header_bold"], True)
check('delete col', (red3.delete_table_column(), tbl3.columns()), (True, 2))
check('delete row', (red3.delete_table_row(), tbl3.rows()), (True, 3))
# header fill survives Word round-trip
tbl3.cellAt(1, 0).firstCursorPosition().insertText("x")
rt3 = os.path.join(tmp, "rt3.docx")
text_to_docx(rt3, "", document=red3.document())
from smoothcursor.documents.docx_io import _docx_cell_fill as _fill
check('fill exported', _fill(pydocx.Document(rt3).tables[0].rows[0].cells[0]),
      "#d6e4f0")
qdoc3 = QTextDocument()
docx_to_document(rt3, qdoc3)
found_fill = None
_b = qdoc3.firstBlock()
while _b.isValid():
    _pp = QTextCursor(qdoc3)
    _pp.setPosition(_b.position())
    _t = _pp.currentTable()
    if _t is not None:
        from smoothcursor.documents.docx_io import _qt_cell_fill_hex as _qfill
        found_fill = _qfill(qdoc3, _t, 0, 0)
        break
    _b = _b.next()
check('fill reimported', found_fill, "#d6e4f0")

# --- custom JSON themes ---
import json as _json
from smoothcursor.core.theme import load_custom_themes as _load, THEMES as _TH
tdir = tempfile.mkdtemp()
with open(os.path.join(tdir, "midnight.json"), "w") as _f:
    _json.dump({"accent": "#ff9e00", "editor_bg": "#101020"}, _f)
with open(os.path.join(tdir, "bad.json"), "w") as _f:
    _f.write("{nope")
loaded, errors = _load(tdir)
check('theme loaded', loaded, ["midnight"])
check('theme bad flagged', "bad.json" in errors, True)
check('theme accent', _TH["midnight"]["accent"], "#ff9e00")
check('theme fallback bg', _TH["midnight"]["editor_bg"], "#101020")

print('fails:', fails)
sys.exit(1 if fails else 0)
